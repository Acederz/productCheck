"""无需填写字段规则：Excel 导入与按大类查询。"""

from datetime import datetime

from openpyxl import load_workbook

from app.extensions import db
from app.models.skip_field_rule import SkipFieldRule, SkipFieldRuleVersion
from app.utils.excel_helper import cell_str

# 中文列名 -> 任务/API 英文字段（可禁用列；不含大类/区隔）
SKIP_FIELD_CN_TO_EN = {
    "类别": "category_type",
    "主材质": "material_main",
    "辅材质": "material_aux",
    "包装方式": "packaging",
    "尺寸": "size",
    "卷数": "roll_count",
    "总入数": "total_count",
}
SKIP_FIELD_EN_TO_CN = {v: k for k, v in SKIP_FIELD_CN_TO_EN.items()}

# 导入时忽略的列名
_SKIP_IMPORT_COLUMNS = {"大类", "区隔"}


class SkipFieldRuleService:
    """无需填写字段规则业务逻辑。"""

    def get_latest_version(self) -> SkipFieldRuleVersion | None:
        """获取最新规则版本（按 id 降序）。"""
        return (
            SkipFieldRuleVersion.query.order_by(SkipFieldRuleVersion.id.desc()).first()
        )

    def get_disabled_fields(self, category_large: str) -> list[str]:
        """按大类返回禁用字段的英文名列表；未命中或未导入时返回空列表。"""
        large = (category_large or "").strip()
        if not large:
            return []

        version = self.get_latest_version()
        if not version:
            return []

        rule = SkipFieldRule.query.filter_by(
            version_id=version.id, category_large=large, is_active=True
        ).first()
        if not rule or not rule.disabled_fields:
            return []

        result = []
        for cn in rule.disabled_fields:
            en = SKIP_FIELD_CN_TO_EN.get(cn)
            if en:
                result.append(en)
        return result

    def import_from_excel(
        self, file_path: str, user_id: int | None, remark: str = ""
    ) -> dict:
        """从 Excel 导入并生成新版本（全量覆盖，旧版本保留作历史）。"""
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.worksheets[0]

        # 解析表头
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            wb.close()
            raise ValueError("Excel 表头为空")

        headers = [cell_str(h) if h is not None else "" for h in header_row]
        if "大类" not in headers:
            wb.close()
            raise ValueError("缺少「大类」列")

        large_col = headers.index("大类")
        # 可禁用列：表头中除大类/区隔外的已知字段列
        field_cols: dict[int, str] = {}
        for idx, name in enumerate(headers):
            if name in _SKIP_IMPORT_COLUMNS or not name:
                continue
            if name in SKIP_FIELD_CN_TO_EN:
                field_cols[idx] = name

        # 按大类聚合禁用字段（同大类多行取并集）
        disabled_by_large: dict[str, set[str]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            large_val = (
                cell_str(row[large_col]) if len(row) > large_col else ""
            )
            if not large_val:
                continue

            disabled: set[str] = set()
            for col_idx, cn_name in field_cols.items():
                if len(row) <= col_idx:
                    continue
                cell_val = cell_str(row[col_idx])
                if cell_val == "无":
                    disabled.add(cn_name)

            if disabled:
                disabled_by_large.setdefault(large_val, set()).update(disabled)

        wb.close()

        version_no = datetime.now().strftime("v%Y%m%d%H%M%S")
        version = SkipFieldRuleVersion(
            version_no=version_no,
            remark=remark or "Excel 导入",
            created_by=user_id,
        )
        db.session.add(version)
        db.session.flush()

        rules = []
        for large, fields in disabled_by_large.items():
            rules.append(
                SkipFieldRule(
                    version_id=version.id,
                    category_large=large,
                    disabled_fields=sorted(fields),
                    is_active=True,
                )
            )
        if rules:
            db.session.bulk_save_objects(rules)

        db.session.commit()

        return {
            "version_no": version_no,
            "rule_count": len(rules),
        }
