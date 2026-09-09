"""无需填写字段规则：导入解析与按大类查询。"""
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from flask_jwt_extended import create_access_token
from openpyxl import Workbook

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app import create_app
from app.constants import ROLE_ADMIN, ROLE_OPERATOR, TASK_STATUS_PENDING
from app.extensions import db
from app.models.log import FieldChangeLog, OperationLog
from app.models.skip_field_rule import SkipFieldRule, SkipFieldRuleVersion
from app.models.task import ClassificationTask, TaskDraft
from app.models.user import User
from app.services.skip_field_rule_service import SkipFieldRuleService
from app.services.task_service import TaskService

UT_PREFIX = "UT_SKIP_"


def _write_xlsx(rows: list[tuple]) -> str:
    """生成测试用 Excel 并返回临时文件路径。"""
    wb = Workbook()
    ws = wb.active
    headers = ["大类", "区隔", "类别", "主材质", "辅材质", "包装方式", "尺寸", "卷数", "总入数"]
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


class SkipFieldRuleServiceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.ctx = cls.app.app_context()
        cls.ctx.push()

    @classmethod
    def tearDownClass(cls):
        cls._cleanup()
        cls.ctx.pop()

    @classmethod
    def _cleanup(cls):
        db.session.rollback()
        versions = SkipFieldRuleVersion.query.filter(
            SkipFieldRuleVersion.version_no.like(f"{UT_PREFIX}%")
        ).all()
        ids = [v.id for v in versions]
        if ids:
            SkipFieldRule.query.filter(SkipFieldRule.version_id.in_(ids)).delete(
                synchronize_session=False
            )
            SkipFieldRuleVersion.query.filter(SkipFieldRuleVersion.id.in_(ids)).delete(
                synchronize_session=False
            )
            db.session.commit()

    def tearDown(self):
        self._cleanup()

    def test_import_and_query_by_large(self):
        path = _write_xlsx(
            [
                ("拖把类", None, None, None, None, None, None, "无", "无"),
                ("手套类", None, None, None, None, None, None, None, "无"),
            ]
        )
        svc = SkipFieldRuleService()
        # 测试里用可控 version_no：实现可接受 remark 含 UT_PREFIX，或测试后改 version_no
        result = svc.import_from_excel(path, user_id=None, remark=f"{UT_PREFIX}t1")
        self.assertGreaterEqual(result["rule_count"], 2)
        # 将最新 version_no 改成 UT 前缀便于清理（若服务已用 v时间戳）
        ver = svc.get_latest_version()
        self.assertIsNotNone(ver)
        ver.version_no = f"{UT_PREFIX}{ver.version_no}"
        db.session.commit()

        self.assertEqual(
            sorted(svc.get_disabled_fields("拖把类")),
            ["roll_count", "total_count"],
        )
        self.assertEqual(svc.get_disabled_fields("手套类"), ["total_count"])
        self.assertEqual(svc.get_disabled_fields("保鲜膜类"), [])
        self.assertEqual(svc.get_disabled_fields(""), [])

    def test_full_replace_on_reimport(self):
        path1 = _write_xlsx([("拖把类", None, None, None, None, None, None, "无", "无")])
        path2 = _write_xlsx([("刷子类", None, None, None, None, None, None, "无", None)])
        svc = SkipFieldRuleService()
        svc.import_from_excel(path1, None, f"{UT_PREFIX}a")
        v1 = svc.get_latest_version()
        v1.version_no = f"{UT_PREFIX}a_{v1.id}"
        db.session.commit()
        svc.import_from_excel(path2, None, f"{UT_PREFIX}b")
        v2 = svc.get_latest_version()
        v2.version_no = f"{UT_PREFIX}b_{v2.id}"
        db.session.commit()
        self.assertEqual(svc.get_disabled_fields("拖把类"), [])
        self.assertEqual(svc.get_disabled_fields("刷子类"), ["roll_count"])

    def test_missing_large_header_raises(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["区隔", "卷数"])
        ws.append(["x", "无"])
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        import os
        os.close(fd)
        wb.save(path)
        with self.assertRaises(ValueError):
            SkipFieldRuleService().import_from_excel(path, None, f"{UT_PREFIX}bad")


class SkipFieldRuleApiTests(unittest.TestCase):
    """无需填写字段规则 HTTP API 冒烟测试。"""

    _UT_OPERATOR = "ut_skip_field_op"

    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.client = cls.app.test_client()
        cls.ctx = cls.app.app_context()
        cls.ctx.push()

    @classmethod
    def tearDownClass(cls):
        SkipFieldRuleServiceTests._cleanup()
        op = User.query.filter_by(username=cls._UT_OPERATOR).first()
        if op:
            db.session.delete(op)
            db.session.commit()
        cls.ctx.pop()

    def tearDown(self):
        SkipFieldRuleServiceTests._cleanup()

    def _admin_token(self) -> str:
        admin = User.query.filter_by(role=ROLE_ADMIN, is_active=True).first()
        self.assertIsNotNone(admin, "需要至少一名启用中的管理员账号")
        return create_access_token(identity=str(admin.id))

    def _operator_token(self) -> str:
        op = User.query.filter_by(username=self._UT_OPERATOR).first()
        if not op:
            op = User(username=self._UT_OPERATOR, role=ROLE_OPERATOR, is_active=True)
            op.set_password("ut_test_pass")
            db.session.add(op)
            db.session.commit()
        return create_access_token(identity=str(op.id))

    def test_disabled_fields_api(self):
        """导入规则后，disabled-fields 接口应返回对应英文字段。"""
        path = _write_xlsx(
            [("拖把类", None, None, None, None, None, None, "无", "无")]
        )
        svc = SkipFieldRuleService()
        svc.import_from_excel(path, user_id=None, remark=f"{UT_PREFIX}api")
        ver = svc.get_latest_version()
        ver.version_no = f"{UT_PREFIX}api_{ver.id}"
        db.session.commit()

        resp = self.client.get(
            "/api/skip-field-rules/disabled-fields?category_large=拖把类",
            headers={"Authorization": f"Bearer {self._operator_token()}"},
        )
        self.assertEqual(resp.status_code, 200)
        body = json.loads(resp.data)
        self.assertEqual(body["code"], 200)
        self.assertEqual(
            set(body["data"]["fields"]),
            {"roll_count", "total_count"},
        )

    def test_import_requires_admin(self):
        """操作员调用 import 接口应返回 403。"""
        xlsx_bytes = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.append(["大类", "卷数", "总入数"])
        ws.append(["拖把类", "无", "无"])
        wb.save(xlsx_bytes)
        xlsx_bytes.seek(0)

        resp = self.client.post(
            "/api/skip-field-rules/import",
            data={"file": (xlsx_bytes, "skip_rules.xlsx")},
            headers={"Authorization": f"Bearer {self._operator_token()}"},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 403)
        body = json.loads(resp.data)
        self.assertEqual(body["code"], 403)
        self.assertIn("管理员", body["message"])

    @patch.object(SkipFieldRuleService, "get_latest_version", return_value=None)
    def test_version_api_empty(self, _mock_get_latest):
        """无导入数据时 version 接口应返回空 data 与提示消息。"""
        resp = self.client.get(
            "/api/skip-field-rules/version",
            headers={"Authorization": f"Bearer {self._operator_token()}"},
        )
        self.assertEqual(resp.status_code, 200)
        body = json.loads(resp.data)
        self.assertEqual(body["code"], 200)
        self.assertIsNone(body["data"])
        self.assertIn("尚未导入无需填写字段规则", body["message"])

    def test_import_as_admin_success(self):
        """管理员 multipart 导入应成功并写入操作日志，disabled-fields 与规则一致。"""
        path = _write_xlsx(
            [("拖把类", None, None, None, None, None, None, "无", "无")]
        )
        admin = User.query.filter_by(role=ROLE_ADMIN, is_active=True).first()
        self.assertIsNotNone(admin, "需要至少一名启用中的管理员账号")

        try:
            with open(path, "rb") as fp:
                resp = self.client.post(
                    "/api/skip-field-rules/import",
                    data={
                        "file": (fp, "skip_rules.xlsx"),
                        "remark": f"{UT_PREFIX}admin_import",
                    },
                    headers={"Authorization": f"Bearer {self._admin_token()}"},
                    content_type="multipart/form-data",
                )
            self.assertEqual(resp.status_code, 200)
            body = json.loads(resp.data)
            self.assertEqual(body["code"], 200)
            self.assertGreaterEqual(body["data"]["rule_count"], 1)
            self.assertIn("导入成功", body["message"])

            svc = SkipFieldRuleService()
            ver = svc.get_latest_version()
            self.assertIsNotNone(ver)
            imported_version_no = ver.version_no
            ver.version_no = f"{UT_PREFIX}{ver.version_no}"
            db.session.commit()

            resp2 = self.client.get(
                "/api/skip-field-rules/disabled-fields?category_large=拖把类",
                headers={"Authorization": f"Bearer {self._operator_token()}"},
            )
            self.assertEqual(resp2.status_code, 200)
            body2 = json.loads(resp2.data)
            self.assertEqual(
                set(body2["data"]["fields"]),
                {"roll_count", "total_count"},
            )

            log = (
                OperationLog.query.filter_by(
                    user_id=admin.id,
                    action="import_skip_field_rules",
                    target_type="skip_field_rule_version",
                    target_id=ver.id,
                )
                .order_by(OperationLog.id.desc())
                .first()
            )
            self.assertIsNotNone(log)
            self.assertGreaterEqual(log.detail_json.get("rule_count", 0), 1)
            self.assertEqual(log.detail_json.get("version_no"), imported_version_no)
        finally:
            if os.path.exists(path):
                os.unlink(path)


class SkipFieldTaskValidationTests(unittest.TestCase):
    """TaskService 暂存/更新/提交与无需填写字段规则联动校验。"""

    _UT_OP = "UT_SKIP_TASK_op"

    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.ctx = cls.app.app_context()
        cls.ctx.push()

    @classmethod
    def tearDownClass(cls):
        cls._cleanup_all()
        cls.ctx.pop()

    @classmethod
    def _cleanup_tasks_for_op(cls, op_id: int):
        task_ids = [
            t.id
            for t in ClassificationTask.query.filter_by(assignee_id=op_id).all()
        ]
        if task_ids:
            FieldChangeLog.query.filter(FieldChangeLog.task_id.in_(task_ids)).delete(
                synchronize_session=False
            )
            TaskDraft.query.filter(TaskDraft.task_id.in_(task_ids)).delete(
                synchronize_session=False
            )
            ClassificationTask.query.filter(
                ClassificationTask.id.in_(task_ids)
            ).delete(synchronize_session=False)
        db.session.commit()

    @classmethod
    def _cleanup_all(cls):
        SkipFieldRuleServiceTests._cleanup()
        op = User.query.filter_by(username=cls._UT_OP).first()
        if op:
            cls._cleanup_tasks_for_op(op.id)
            db.session.delete(op)
            db.session.commit()

    def setUp(self):
        SkipFieldRuleServiceTests._cleanup()
        path = _write_xlsx(
            [
                ("拖把类", None, None, None, None, None, None, "无", "无"),
                ("湿巾类", None, "无", None, None, None, None, None, None),
            ]
        )
        svc = SkipFieldRuleService()
        svc.import_from_excel(path, user_id=None, remark=f"{UT_PREFIX}task_val")
        ver = svc.get_latest_version()
        ver.version_no = f"{UT_PREFIX}task_val_{ver.id}"
        db.session.commit()
        if os.path.exists(path):
            os.unlink(path)

        op = User.query.filter_by(username=self._UT_OP).first()
        if not op:
            op = User(username=self._UT_OP, role=ROLE_OPERATOR, is_active=True)
            op.set_password("ut_pass")
            db.session.add(op)
            db.session.commit()
        self.op = op
        self.op_id = op.id

        task = ClassificationTask(
            product_id=f"P_skip_{self.op_id}",
            product_name="测试商品",
            platform="淘宝",
            status=TASK_STATUS_PENDING,
            assignee_id=self.op_id,
        )
        db.session.add(task)
        db.session.commit()
        self.task_id = task.id
        self.task = task

    def tearDown(self):
        SkipFieldRuleServiceTests._cleanup()
        self._cleanup_tasks_for_op(self.op_id)

    def test_save_draft_rejects_disabled_value(self):
        svc = TaskService()
        with self.assertRaises(ValueError) as ctx:
            svc.save_draft(
                self.task_id,
                self.op_id,
                {"category_large": "拖把类", "roll_count": "2卷"},
            )
        self.assertIn("卷数", str(ctx.exception))

    def test_submit_skips_required_when_category_type_disabled(self):
        self.task.is_operating = "是"
        self.task.category_large = "湿巾类"
        self.task.category_segment = ["家用"]
        self.task.category_type = None
        db.session.commit()

        svc = TaskService()
        self.assertIsNone(svc._validate_for_submit(self.task))

    def test_update_task_rejects_disabled_value(self):
        self.task.category_large = "拖把类"
        db.session.commit()

        svc = TaskService()
        with self.assertRaises(ValueError) as ctx:
            svc.update_task(
                self.task,
                {"total_count": "100"},
                self.op,
            )
        self.assertIn("总入数", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
